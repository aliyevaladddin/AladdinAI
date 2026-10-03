"""Regression tests for DOCX-to-WRT inline image ordering."""

import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from app.services.wrt_engine_service import docx_to_wrt_sync, wrt_to_docx_sync


def _png_bytes() -> bytes:
    image = Image.new("RGB", (20, 20), "blue")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def test_inline_image_stays_between_surrounding_text_runs():
    doc = Document()
    paragraph = doc.add_paragraph()

    paragraph.add_run("before ")
    paragraph.add_run().add_picture(io.BytesIO(_png_bytes()), width=Inches(0.25))
    paragraph.add_run(" after")

    buf = io.BytesIO()
    doc.save(buf)

    wrt = docx_to_wrt_sync(buf.getvalue())

    before = wrt.index("before")
    image = wrt.index("[img ")
    after = wrt.index("after")

    assert before < image < after


def test_inline_image_is_not_hoisted_before_paragraph_text():
    doc = Document()
    paragraph = doc.add_paragraph()

    paragraph.add_run("prefix ")
    paragraph.add_run().add_picture(io.BytesIO(_png_bytes()), width=Inches(0.25))
    paragraph.add_run(" suffix")

    buf = io.BytesIO()
    doc.save(buf)

    wrt = docx_to_wrt_sync(buf.getvalue())

    assert not wrt.lstrip().startswith("[img ")
    assert wrt.index("prefix") < wrt.index("[img ")
    assert wrt.index("[img ") < wrt.index("suffix")


def _paragraph_run_sequence(paragraph) -> list:
    """Ordered flat sequence of a paragraph's runs: text runs as str, image runs as 'IMG'."""
    seq = []
    for run in paragraph._element.findall(qn("w:r")):
        if run.find(".//" + qn("w:drawing")) is not None:
            seq.append("IMG")
        else:
            t = run.find(qn("w:t"))
            if t is not None and t.text:
                seq.append(t.text)
    return seq


def test_inline_image_round_trips_back_to_docx():
    """DOCX -> WRT -> DOCX: an inline image between text runs survives as a drawing in place."""
    doc = Document()
    paragraph = doc.add_paragraph()

    paragraph.add_run("before ")
    paragraph.add_run().add_picture(io.BytesIO(_png_bytes()), width=Inches(0.25))
    paragraph.add_run(" after")

    buf = io.BytesIO()
    doc.save(buf)

    wrt = docx_to_wrt_sync(buf.getvalue())
    docx2 = wrt_to_docx_sync(wrt)

    doc2 = Document(io.BytesIO(docx2))
    target = None
    for p in doc2.paragraphs:
        if "before" in p.text and "after" in p.text:
            target = p
            break
    assert target is not None, "text and image were split into separate paragraphs"

    # Image restored as a real inline drawing, not literal [img ...] text
    assert "[img" not in target.text
    seq = _paragraph_run_sequence(target)
    assert seq == ["before ", "IMG", " after"]

    # Image bytes embedded in the package
    with zipfile.ZipFile(io.BytesIO(docx2)) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert media, "image bytes missing from docx package"


def test_inline_image_invalid_base64_falls_back_to_inline_placeholder():
    """A malformed inline image between text runs degrades to an inline text placeholder."""
    wrt = 'before [img src="data:image/png;base64,INVALID_BASE64_DATA!!!" alt="broken"] after'
    docx_bytes = wrt_to_docx_sync(wrt)

    doc = Document(io.BytesIO(docx_bytes))
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "before [Image: broken] after" in texts, texts
